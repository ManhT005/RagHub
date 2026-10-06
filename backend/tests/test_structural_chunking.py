import io
from uuid import uuid4

import docx
from raghub_core.domain.ingestion.chunker import chunk_sections

from app.infrastructure.parsing.docx import parse_docx
from app.infrastructure.parsing.html import parse_html
from app.infrastructure.parsing.xlsx import parse_xlsx
from tests.test_document_parsers import _xlsx


def test_docx_heading_path_and_short_paragraph_merging():
    document = docx.Document()
    document.add_heading("Guide", level=1)
    document.add_heading("Setup", level=2)
    for i in range(30):
        document.add_paragraph(f"Step {i}: configure the service.")
    buffer = io.BytesIO()
    document.save(buffer)
    blocks = parse_docx(buffer.getvalue())
    assert blocks[-1].heading_path == ("Guide", "Setup")
    chunks = chunk_sections(blocks, uuid4())
    assert len(chunks) < 10
    assert all(chunk.heading == "Guide / Setup" for chunk in chunks)
    assert all(any(f"Step {i}:" in c.content for c in chunks) for i in range(30))


def test_xlsx_chunks_repeat_header_and_preserve_every_row():
    blocks = parse_xlsx(
        _xlsx({"Data": [["Name", "Value"], *[[f"record-{i}", "value " * 25] for i in range(100)]]})
    )
    chunks = chunk_sections(blocks, uuid4(), target_tokens=120, overlap_tokens=12)
    assert len(chunks) > 1
    assert all("| Name | Value |" in chunk.content for chunk in chunks)
    assert all(chunk.token_count <= 120 for chunk in chunks)
    for i in range(100):
        assert sum(f"| record-{i} |" in c.content for c in chunks) == 1


def test_html_uses_heading_tags_and_removes_nested_hidden_subtrees():
    blocks = parse_html(
        b"<h1>Guide</h1><p>Short label</p><h2>Setup</h2>"
        b"<div hidden><b>secret</b>still secret</div><p>Visible</p>"
    )
    assert blocks[1].type == "paragraph"
    assert blocks[1].heading_path == ("Guide",)
    assert blocks[-1].heading_path == ("Guide", "Setup")
    assert "secret" not in " ".join(b.content for b in blocks)
