"""XLSX adapter: read-only + data-only via openpyxl, no macro/formula execution."""

from __future__ import annotations

import io

import openpyxl
from raghub_core.domain.ingestion.limits import (
    MAX_POPULATED_CELLS,
    MAX_SHEETS,
    check_compressed_size,
    check_signature,
    check_zip_container,
)
from raghub_core.domain.ingestion.parser import (
    DocumentLimitError,
    EmptyExtractedTextError,
    ParsedBlock,
    ParsedSection,
)


def parse_xlsx(content: bytes, source_name: str = "document.xlsx") -> list[ParsedSection]:
    check_compressed_size(len(content))
    check_signature(content, ".xlsx")
    check_zip_container(content, ".xlsx")
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise EmptyExtractedTextError("The XLSX file cannot be opened.") from exc
    try:
        names = workbook.sheetnames
        if len(names) > MAX_SHEETS:
            raise DocumentLimitError(f"Workbook exceeds {MAX_SHEETS} sheets.")
        sections: list[ParsedSection] = []
        cells = 0
        for sheet in names:
            view = workbook[sheet]
            rows: list[str] = []
            row_numbers = []
            min_row, max_row = view.min_row, view.max_row
            min_col, max_col = view.min_column, view.max_column
            for number, row in enumerate(
                view.iter_rows(
                    min_row=min_row,
                    max_row=max_row,
                    min_col=min_col,
                    max_col=max_col,
                    values_only=True,
                ),
                start=min_row,
            ):
                values = [
                    "" if value is None else str(value).strip().replace("|", "/").replace("\n", " ")
                    for value in row
                ]
                while values and not values[-1]:
                    values.pop()
                if any(values):
                    cells += sum(1 for value in values if value)
                    if cells > MAX_POPULATED_CELLS:
                        raise DocumentLimitError(
                            f"Workbook exceeds {MAX_POPULATED_CELLS} populated cells."
                        )
                    rows.append("| " + " | ".join(values) + " |")
                    row_numbers.append(number)
            if rows:
                width = max(row.count("|") - 1 for row in rows)
                table = [rows[0], "| " + " | ".join(["---"] * width) + " |", *rows[1:]]
                sections.append(
                    ParsedBlock(
                        "\n".join(table),
                        source_name,
                        len(sections),
                        heading=f"Sheet: {sheet} (rows {min_row}-{max_row})",
                        type="table",
                        heading_path=(sheet,),
                        metadata={
                            "sheet": sheet,
                            "row_start": row_numbers[0],
                            "row_end": row_numbers[-1],
                            "row_numbers": row_numbers[1:],
                            "header": rows[0],
                            "header_row": row_numbers[0],
                        },
                    )
                )
        if not sections:
            raise EmptyExtractedTextError("The XLSX file contains no data.")
        return sections
    finally:
        workbook.close()
